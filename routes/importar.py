import os
import tempfile
import openpyxl
from datetime import datetime
from flask import Blueprint, render_template, request, redirect, url_for, flash, session
from flask_login import login_required, current_user
from models import db, Processo

importar_bp = Blueprint('importar', __name__)

# Campos obrigatórios e seus labels legíveis
CAMPOS_OBRIGATORIOS = {
    'nome':        'Nome do Servidor',
    'tipo_benef':  'Tipo de Beneficiário',
    'cargo':       'Cargo / Função',
    'orgao':       'Órgão / Secretaria',
    'origem':      'Município de Origem',
    'destino':     'Município de Destino',
    'dt_saida':    'Data de Saída',
    'hr_saida':    'Hora de Saída',
    'dt_chegada':  'Data de Chegada',
    'hr_chegada':  'Hora de Chegada',
    'tipo_viagem': 'Tipo de Viagem',
    'total_diaria':'Total da Diária (R$)',
}

CAMPOS_RECOMENDADOS = {
    'cpf':      'CPF',
    'matricula':'Matrícula',
    'banco':    'Banco',
    'agencia':  'Agência',
    'conta':    'Conta Corrente',
    'un_orc':   'Unidade Orçamentária',
    'proj_atv': 'Projeto / Atividade',
    'elem_desp':'Elemento de Despesa',
}


def _val(sheet, coord, default=''):
    import re
    v = sheet[coord].value
    if v is None:
        return default
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if hasattr(v, 'total_seconds'):
        h = int(v.total_seconds() // 3600)
        m = int((v.total_seconds() % 3600) // 60)
        return f'{h}h {m}min'
    s = str(v).strip()
    if re.match(r'^\d+ ?[-–]', s):
        return default
    return s


def _flt(sheet, coord):
    v = sheet[coord].value
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _int(sheet, coord):
    v = sheet[coord].value
    try:
        return int(v)
    except (TypeError, ValueError):
        return 0


def _time(sheet, coord):
    v = sheet[coord].value
    if v is None:
        return ''
    if hasattr(v, 'strftime'):
        return v.strftime('%H:%M')
    return str(v).strip()


def _tipo_viagem(valor):
    if not valor:
        return ''
    return 'fora_estado' if 'fora do estado' in str(valor).lower() else 'no_estado'


def _sim_nao(valor):
    return 'sim' if str(valor or '').strip().lower() == 'sim' else 'nao'


def _proximo_num_form():
    ano = datetime.utcnow().year
    ultimo = (Processo.query
              .filter(Processo.num_form.like(f'%/{ano}'))
              .order_by(Processo.id.desc())
              .first())
    seq = int(ultimo.num_form.split('/')[0]) + 1 if ultimo else 1
    return f'{str(seq).zfill(3)}/{ano}'


def _parse_xlsx(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    if 'DIÁRIA' not in wb.sheetnames:
        raise ValueError("Aba 'DIÁRIA' não encontrada.")
    d = wb['DIÁRIA']

    STATUS_MAP = {
        'agente': 'Agente Político', 'comissionado': 'Comissionado',
        'efetivo': 'Efetivo', 'contratado': 'Contratado', 'convidado': 'Convidado',
    }
    status_raw = _val(d, 'K6', '')
    status_form = STATUS_MAP.get(status_raw.lower(), status_raw)

    num_form_xlsx = _val(d, 'C7')
    if num_form_xlsx and Processo.query.filter_by(num_form=num_form_xlsx).first():
        num_form = _proximo_num_form()
    elif num_form_xlsx:
        num_form = num_form_xlsx
    else:
        num_form = _proximo_num_form()

    dados = {
        'num_form':     num_form,
        'data_emissao': _val(d, 'H7') or datetime.utcnow().strftime('%Y-%m-%d'),
        'status_form':  status_form,
        'nome':         _val(d, 'A12'),
        'matricula':    str(_val(d, 'G16') or '').strip(),
        'cpf':          '',
        'rg':           '',
        'tipo_benef':   _val(d, 'E14'),
        'cargo':        _val(d, 'A14'),
        'orgao':        _val(d, 'A16'),
        'fund_legal':   _val(d, 'A21'),
        'descricao':    _val(d, 'A24'),
        'origem':       _val(d, 'A32'),
        'destino':      _val(d, 'B32'),
        'dt_saida':     _val(d, 'D32'),
        'hr_saida':     _time(d, 'E32'),
        'dt_chegada':   _val(d, 'F32'),
        'hr_chegada':   _time(d, 'G32'),
        'horas_viagem': _val(d, 'H32'),
        'tipo_viagem':  _tipo_viagem(_val(d, 'A37')),
        'desl_24h':     _sim_nao(_val(d, 'A39')),
        'hosp_outra':   _sim_nao(_val(d, 'E39')),
        'alim_outra':   _sim_nao(_val(d, 'G39')),
        'r_classe':     _val(d, 'A42'),
        'r_tipo':       _val(d, 'B42'),
        'r_qtd_int':    _int(d, 'E42'),
        'r_qtd_meio':   _int(d, 'F42'),
        'r_valor':      _flt(d, 'G42'),
        'total_diaria': _flt(d, 'H42') or _flt(d, 'A44'),
        'un_orc':       str(_val(d, 'D49') or '').strip(),
        'proj_atv':     _val(d, 'E49'),
        'elem_desp':    _val(d, 'F49'),
        'conta':        _val(d, 'C57'),
        'agencia':      str(_val(d, 'E57') or '').strip(),
        'banco':        '',
    }

    if 'RELATÓRIO DE DESPESAS' in wb.sheetnames:
        r = wb['RELATÓRIO DE DESPESAS']
        dados['d_alim']    = _flt(r, 'D10')
        dados['d_transp']  = _flt(r, 'D11')
        dados['d_taxa']    = _flt(r, 'D12')
        dados['d_reparos'] = _flt(r, 'D13')
        dados['d_outros']  = _flt(r, 'D14')

    return dados


def _validar(dados):
    """Retorna (faltam_obrigatorios, faltam_recomendados) como listas de labels."""
    faltam_obrig = [label for campo, label in CAMPOS_OBRIGATORIOS.items()
                    if not str(dados.get(campo, '') or '').strip()
                    or (campo == 'total_diaria' and not float(dados.get(campo, 0) or 0))]
    faltam_rec   = [label for campo, label in CAMPOS_RECOMENDADOS.items()
                    if not str(dados.get(campo, '') or '').strip()]
    return faltam_obrig, faltam_rec


# ── Rotas ────────────────────────────────────────────────────────────────────

@importar_bp.route('/importar', methods=['GET', 'POST'])
@login_required
def importar():
    """Tela 1: seleção de arquivos."""
    if request.method == 'GET':
        return render_template('importar.html', resultados=None)

    arquivos = request.files.getlist('arquivos')
    validos  = [a for a in arquivos if a.filename.lower().endswith('.xlsx')]
    invalidos = [a.filename for a in arquivos
                 if a.filename and not a.filename.lower().endswith('.xlsx')]

    if not validos:
        flash('Nenhum arquivo .xlsx selecionado.', 'erro')
        return redirect(url_for('importar.importar'))

    previsoes = []   # lista de dicts: dados parseados + metadados de validação

    for arquivo in validos:
        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        try:
            arquivo.save(tmp.name)
            tmp.close()
            dados = _parse_xlsx(tmp.name)
            faltam_obrig, faltam_rec = _validar(dados)
            previsoes.append({
                'nome_arquivo':   arquivo.filename,
                'dados':          dados,
                'faltam_obrig':   faltam_obrig,
                'faltam_rec':     faltam_rec,
                'parse_erro':     None,
            })
        except Exception as e:
            previsoes.append({
                'nome_arquivo': arquivo.filename,
                'dados':        None,
                'faltam_obrig': [],
                'faltam_rec':   [],
                'parse_erro':   str(e),
            })
        finally:
            try:
                os.unlink(tmp.name)
            except Exception:
                pass

    # salva previsões na sessão para a tela de revisão acessar
    session['importar_previsoes'] = previsoes
    session['importar_invalidos'] = invalidos

    return redirect(url_for('importar.revisar'))


@importar_bp.route('/importar/revisar', methods=['GET', 'POST'])
@login_required
def revisar():
    """Tela 2: revisão editável — um formulário por arquivo."""
    if request.method == 'GET':
        previsoes = session.get('importar_previsoes')
        if not previsoes:
            flash('Nenhum arquivo em revisão. Faça o upload novamente.', 'erro')
            return redirect(url_for('importar.importar'))
        invalidos = session.get('importar_invalidos', [])
        campos_obrig = CAMPOS_OBRIGATORIOS
        campos_rec   = CAMPOS_RECOMENDADOS
        return render_template('importar_revisar.html',
                               previsoes=previsoes,
                               invalidos=invalidos,
                               campos_obrig=campos_obrig,
                               campos_rec=campos_rec)

    # POST: confirmar e salvar
    previsoes = session.get('importar_previsoes', [])
    resultados = []

    for i, prev in enumerate(previsoes):
        if prev.get('parse_erro'):
            resultados.append({'nome': prev['nome_arquivo'], 'status': 'erro',
                               'mensagem': prev['parse_erro'], 'num_form': None})
            continue

        # coleta dados editados do form — prefixo "p{i}_"
        prefix = f'p{i}_'
        dados = dict(prev['dados'])   # começa com o que foi parseado
        for campo in list(CAMPOS_OBRIGATORIOS.keys()) + list(CAMPOS_RECOMENDADOS.keys()) + [
            'data_emissao','status_form','rg','fund_legal','descricao',
            'horas_viagem','desl_24h','hosp_outra','alim_outra',
            'r_classe','r_tipo','un_orc','proj_atv','elem_desp','conta','agencia','banco',
        ]:
            key = prefix + campo
            if key in request.form:
                dados[campo] = request.form[key].strip()

        # campos numéricos
        for campo in ['r_qtd_int','r_qtd_meio']:
            key = prefix + campo
            if key in request.form:
                try:
                    dados[campo] = int(request.form[key])
                except (ValueError, TypeError):
                    dados[campo] = 0
        for campo in ['r_valor','total_diaria','d_alim','d_transp','d_taxa','d_reparos','d_outros']:
            key = prefix + campo
            if key in request.form:
                try:
                    dados[campo] = float(request.form[key].replace(',', '.'))
                except (ValueError, TypeError):
                    dados[campo] = 0.0

        # re-valida com dados editados
        faltam_obrig, _ = _validar(dados)
        if faltam_obrig:
            resultados.append({'nome': prev['nome_arquivo'], 'status': 'erro',
                               'mensagem': f'Campos obrigatórios em falta: {", ".join(faltam_obrig)}',
                               'num_form': None})
            continue

        try:
            p = Processo(usuario_id=current_user.id, status='rascunho')
            for campo, valor in dados.items():
                if hasattr(p, campo):
                    setattr(p, campo, valor)
            db.session.add(p)
            db.session.flush()
            resultados.append({'nome': prev['nome_arquivo'], 'status': 'ok',
                               'mensagem': f'Salvo como {p.num_form}.', 'num_form': p.num_form})
        except Exception as e:
            db.session.rollback()
            resultados.append({'nome': prev['nome_arquivo'], 'status': 'erro',
                               'mensagem': f'Erro ao salvar: {str(e)}', 'num_form': None})

    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(f'Erro ao salvar no banco: {str(e)}', 'erro')
        return redirect(url_for('importar.revisar'))

    session.pop('importar_previsoes', None)
    session.pop('importar_invalidos', None)

    ok  = sum(1 for r in resultados if r['status'] == 'ok')
    err = sum(1 for r in resultados if r['status'] == 'erro')
    flash(f'{ok} processo(s) importado(s) com sucesso. {err} com erro.', 'ok' if err == 0 else 'erro')
    return render_template('importar.html', resultados=resultados)
